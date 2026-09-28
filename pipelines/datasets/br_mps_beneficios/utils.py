"""Pure download and cleaning functions for br_mps_beneficios.

No Prefect imports: the one-shot bootstrap under ``models/br_mps_beneficios/code/``
imports these same functions, so the transform exists in one place only.

The source ships the same logical table in three incompatible layouts, and the
reader normalises all of them to one frame before aggregation:

===  ==========================  =====================  ====================
era  package                     layout                 espécie
===  ==========================  =====================  ====================
V1   annual zips 2012-2018       19-col ;-CSV, latin-1  label only
V2   monthly CSV Dec/18-May/23    13-col ;-CSV, latin-1  label only
V3   monthly XLSX Jun/23-        27-col, 2-row header   code + label
===  ==========================  =====================  ====================

``mantidos`` is a fourth layout: fixed-width-padded ;-CSV with no competência
column (it comes from the filename), no espécie code, and labels truncated to
20 characters.
"""

from __future__ import annotations

import csv
import re
import unicodedata
import zipfile
from datetime import date
from pathlib import Path

import pandas as pd
import pyarrow as pa
import pyarrow.parquet as pq

from pipelines.datasets.br_mps_beneficios.constants import constants

MESES_PT = {
    "janeiro": 1,
    "fevereiro": 2,
    "marco": 3,
    "março": 3,
    "abril": 4,
    "maio": 5,
    "junho": 6,
    "julho": 7,
    "agosto": 8,
    "setembro": 9,
    "outubro": 10,
    "novembro": 11,
    "dezembro": 12,
}


# --------------------------------------------------------------------------
# normalisation helpers
# --------------------------------------------------------------------------
def strip_accents(s: str) -> str:
    return (
        unicodedata.normalize("NFKD", str(s))
        .encode("ascii", "ignore")
        .decode()
    )


def norm_token(s: str) -> str:
    """Lowercase, unaccent, collapse to single-spaced alphanumeric tokens."""
    return re.sub(
        r"\s+", " ", re.sub(r"[^a-z0-9]", " ", strip_accents(s).lower())
    ).strip()


def norm_compact(s: str) -> str:
    """Strip every non-alphanumeric character.

    Needed because the source and the directory disagree on apostrophes and
    hyphens in municipality names: SUIBE writes "Santana do Livramento" where
    IBGE writes "Sant'Ana do Livramento", and the same applies to the
    "d'Água" / "d'Ávila" family. Collapsing to letters and digits makes the
    two forms agree without a per-name exception.
    """
    return re.sub(r"[^a-z0-9]", "", strip_accents(s).lower())


def expand_abbrev(s: str) -> str:
    exp = constants.ABREV_EXPANSION.value
    return " ".join(exp.get(t, t) for t in norm_token(s).split())


# --------------------------------------------------------------------------
# reference tables
# --------------------------------------------------------------------------
def especie_label_index() -> dict[str, int]:
    """Map every spelling of an espécie label to its code.

    Built from the official dictionary under both plain normalisation and
    abbreviation expansion, then the hand-checked alias list.
    """
    idx: dict[str, int] = {}
    for code, label in constants.ESPECIE.value.items():
        idx.setdefault(norm_token(label), code)
        idx.setdefault(expand_abbrev(label), code)
    idx.update(constants.ESPECIE_LABEL_ALIAS.value)
    for key, code in list(constants.ESPECIE_LABEL_ALIAS.value.items()):
        idx.setdefault(expand_abbrev(key), code)
    return idx


def categoria_index() -> dict[int, str]:
    return {
        c: cat
        for cat, codes in constants.CATEGORIA.value.items()
        for c in codes
    }


def municipio_index() -> dict[tuple[str, str], str]:
    """(sigla_uf, key) -> id_municipio, under both normalisation forms.

    Each name is registered under ``norm_token`` and under ``norm_compact`` so
    that punctuation differences resolve without an exception entry. A compact
    key that would collide with a different municipality in the same UF is
    dropped rather than allowed to resolve ambiguously.
    """
    idx: dict[tuple[str, str], str] = {}
    compact: dict[tuple[str, str], set[str]] = {}

    def add(uf: str, nome: str, id_mun: str) -> None:
        idx[(uf, norm_token(nome))] = id_mun
        compact.setdefault((uf, norm_compact(nome)), set()).add(id_mun)

    with open(constants.MUNICIPIO_DIRECTORY.value, encoding="utf-8") as f:
        for r in csv.DictReader(f):
            add(r["sigla_uf"], r["nome"], r["id_municipio"])
    with open(constants.MUNICIPIO_CROSSWALK.value, encoding="utf-8") as f:
        for r in csv.DictReader(f):
            add(r["sigla_uf"], r["nome_suibe"], r["id_municipio"])

    for key, ids in compact.items():
        if len(ids) == 1:
            idx.setdefault(key, next(iter(ids)))
    return idx


def lookup_municipio(
    idx: dict[tuple[str, str], str], uf: str, nome: str
) -> str | None:
    return idx.get((uf, norm_token(nome))) or idx.get((uf, norm_compact(nome)))


def salario_minimo_for(competencia: int) -> float:
    """Nominal minimum wage in force in a YYYYMM competência."""
    applicable = [
        k for k in constants.SALARIO_MINIMO.value if k <= competencia
    ]
    if not applicable:
        raise ValueError(
            f"no minimum wage on record at or before {competencia}"
        )
    return constants.SALARIO_MINIMO.value[max(applicable)]


def validate_reference_tables() -> None:
    """Fail fast if the espécie/categoria tables drift out of agreement."""
    especies = set(constants.ESPECIE.value)
    mapped = [c for codes in constants.CATEGORIA.value.values() for c in codes]
    missing = especies - set(mapped)
    extra = set(mapped) - especies
    dupes = {c for c in mapped if mapped.count(c) > 1}
    if missing or extra or dupes:
        raise ValueError(
            f"CATEGORIA inconsistent: missing={sorted(missing)} "
            f"extra={sorted(extra)} duplicated={sorted(dupes)}"
        )
    if len(constants.FAIXA_ETARIA_BINS.value) - 1 != len(
        constants.FAIXA_ETARIA_LABELS.value
    ):
        raise ValueError("FAIXA_ETARIA bins and labels disagree")


# --------------------------------------------------------------------------
# field parsers
# --------------------------------------------------------------------------
def parse_competencia(raw) -> int | None:
    """Normalise every competência spelling the source uses to YYYYMM."""
    if raw is None or (isinstance(raw, float) and pd.isna(raw)):
        return None
    if isinstance(raw, (date,)):
        return raw.year * 100 + raw.month
    s = str(raw).strip()
    if re.fullmatch(r"\d{6}", s):  # V3: 202512
        return int(s)
    m = re.fullmatch(r"(\d{4})-(\d{2})-\d{2}", s)  # V2: 2018-12-01
    if m:
        return int(m.group(1)) * 100 + int(m.group(2))
    m = re.fullmatch(r"([A-Za-zçÇ]+)/(\d{4})", s)  # V1: janeiro/2017
    if m:
        mes = MESES_PT.get(strip_accents(m.group(1)).lower())
        if mes:
            return int(m.group(2)) * 100 + mes
    m = re.fullmatch(r"(\d{4})-(\d{2})", s)
    if m:
        return int(m.group(1)) * 100 + int(m.group(2))
    return None


def parse_data(raw) -> date | None:
    """Parse a birth date; the source writes absent dates as 00/00/0000."""
    if raw is None or (isinstance(raw, float) and pd.isna(raw)):
        return None
    if isinstance(raw, date):
        return raw
    s = str(raw).strip()
    if not s or s.startswith("00/00"):
        return None
    for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d/%m/%y"):
        try:
            return pd.to_datetime(s, format=fmt).date()
        except (ValueError, TypeError):
            continue
    return None


def parse_decimal(raw) -> float | None:
    """Parse Qt SM RMI / Vl MR, which use pt-BR decimal notation in the CSVs."""
    if raw is None or (isinstance(raw, float) and pd.isna(raw)):
        return None
    if isinstance(raw, (int, float)):
        return float(raw)
    s = str(raw).strip()
    if not s:
        return None
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return None


def parse_mun_resid(raw) -> tuple[str | None, str | None]:
    """Split ``Mun Resid`` into (sigla_uf, municipality name).

    The field is documented as "Código da Gerência-Executiva (GEX) seguido da
    sigla da UF e do Município de Residência", so the leading digits are an
    INSS administrative region and carry no municipal identity. The separate
    ``UF`` column is the *granting agency's* UF and disagrees with this one, so
    it must never be used for geography.
    """
    if raw is None:
        return None, None
    s = str(raw).strip()
    if s in constants.MUN_SENTINELS.value or "Zerada" in s:
        return None, None
    m = re.match(r"^\d+\s*-\s*([A-Za-z]{2})\s*-\s*(.+)$", s)
    if not m:
        return None, None
    nome = m.group(2).strip()
    return m.group(1).upper(), (nome or None)


def idade_em(nascimento: date | None, competencia: int) -> int | None:
    if nascimento is None:
        return None
    ano, mes = divmod(competencia, 100)
    idade = ano - nascimento.year - ((mes, 1) < (nascimento.month, 1))
    return idade if 0 <= idade <= 130 else None


def faixa_etaria(idade: int | None) -> str | None:
    if idade is None:
        return None
    bins = constants.FAIXA_ETARIA_BINS.value
    labels = constants.FAIXA_ETARIA_LABELS.value
    for i in range(len(labels)):
        if bins[i] <= idade < bins[i + 1]:
            return labels[i]
    return None


def clean_clientela(raw) -> str | None:
    n = norm_token(raw or "")
    if n.startswith("urban"):
        return "urbana"
    if n.startswith("rural"):
        return "rural"
    return None


def clean_sexo(raw) -> str | None:
    n = norm_token(raw or "")
    if n.startswith("masc"):
        return "masculino"
    if n.startswith("fem"):
        return "feminino"
    return None


# --------------------------------------------------------------------------
# readers — one normalised frame out of three source layouts
# --------------------------------------------------------------------------
CONCEDIDO_FIELDS = [
    "competencia",
    "especie_codigo",
    "especie_label",
    "dt_nascimento",
    "sexo",
    "clientela",
    "mun_resid",
    "qt_sm_rmi",
]


def _col(header: list[str], *wanted: str) -> int | None:
    """Locate a column by normalised name, tolerating the source's drift."""
    norm = [norm_token(h or "") for h in header]
    for w in wanted:
        wn = norm_token(w)
        for i, h in enumerate(norm):
            if h == wn:
                return i
    for w in wanted:
        wn = norm_token(w)
        for i, h in enumerate(norm):
            if h.startswith(wn) or (wn.startswith(h) and h):
                return i
    return None


def iter_concedido_csv(path: Path, chunk: int = 500_000):
    """Stream a V1/V2 semicolon CSV (latin-1) as normalised record dicts."""
    with open(path, encoding="latin-1", errors="replace", newline="") as f:
        header = next(csv.reader(f, delimiter=";"))
        i_comp = _col(header, "Competência concessão")
        i_esp = _col(header, "Espécie")
        i_nasc = _col(header, "Dt Nascimento", "Data Nascimento")
        i_sexo = _col(header, "Sexo.", "Sexo")
        i_cli = _col(header, "Clientela")
        i_mun = _col(header, "Mun Resid", "Município")
        i_rmi = _col(header, "Qt SM RMI")
        missing = [
            n
            for n, i in [
                ("competência", i_comp),
                ("espécie", i_esp),
                ("nascimento", i_nasc),
                ("sexo", i_sexo),
                ("clientela", i_cli),
                ("município", i_mun),
                ("Qt SM RMI", i_rmi),
            ]
            if i is None
        ]
        if missing:
            raise ValueError(
                f"{path.name}: columns not found: {missing} in {header}"
            )
        buf = []
        for row in csv.reader(f, delimiter=";"):
            if len(row) <= max(
                i_comp, i_esp, i_nasc, i_sexo, i_cli, i_mun, i_rmi
            ):
                continue
            buf.append(
                {
                    "competencia": row[i_comp],
                    "especie_codigo": None,
                    "especie_label": row[i_esp],
                    "dt_nascimento": row[i_nasc],
                    "sexo": row[i_sexo],
                    "clientela": row[i_cli],
                    "mun_resid": row[i_mun],
                    "qt_sm_rmi": row[i_rmi],
                }
            )
            if len(buf) >= chunk:
                yield buf
                buf = []
        if buf:
            yield buf


def iter_concedido_xlsx(path: Path, chunk: int = 500_000):
    """Stream a V3 workbook. Row 0 is a title banner, row 1 the real header,
    and ``Espécie``/``CID``/``Despacho`` each occupy a code+label column pair."""
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = ws.iter_rows(values_only=True)
    first = next(rows)
    header = list(next(rows))
    # the banner row is absent in a handful of months
    if sum(1 for c in first if c) > 3:
        header, first = list(first), None
    i_comp = _col(header, "Competência concessão")
    i_esp = _col(header, "Espécie")
    i_nasc = _col(header, "Dt Nascimento", "Data Nascimento")
    i_sexo = _col(header, "Sexo.", "Sexo")
    i_cli = _col(header, "Clientela")
    i_mun = _col(header, "Mun Resid")
    i_rmi = _col(header, "Qt SM RMI")
    if i_esp is None or i_mun is None:
        raise ValueError(f"{path.name}: header not recognised: {header}")
    # Espécie is a (code, label) pair; _col finds the first of the two.
    i_esp_lab = (
        i_esp + 1
        if i_esp + 1 < len(header)
        and norm_token(header[i_esp + 1] or "") == norm_token("Espécie")
        else i_esp
    )
    buf = []
    for row in rows:
        if row is None or all(c is None for c in row):
            continue
        buf.append(
            {
                "competencia": row[i_comp] if i_comp is not None else None,
                "especie_codigo": row[i_esp] if i_esp_lab != i_esp else None,
                "especie_label": row[i_esp_lab],
                "dt_nascimento": row[i_nasc] if i_nasc is not None else None,
                "sexo": row[i_sexo] if i_sexo is not None else None,
                "clientela": row[i_cli] if i_cli is not None else None,
                "mun_resid": row[i_mun],
                "qt_sm_rmi": row[i_rmi] if i_rmi is not None else None,
            }
        )
        if len(buf) >= chunk:
            yield buf
            buf = []
    if buf:
        yield buf
    wb.close()


def iter_concedido(path: Path, chunk: int = 500_000):
    if path.suffix.lower() in (".xlsx", ".xls"):
        yield from iter_concedido_xlsx(path, chunk)
    elif path.suffix.lower() == ".zip":
        with zipfile.ZipFile(path) as z:
            inner = z.namelist()[0]
            tmp = path.parent / inner
            if not tmp.exists():
                z.extract(inner, path.parent)
            try:
                yield from iter_concedido_csv(tmp, chunk)
            finally:
                tmp.unlink(missing_ok=True)
    else:
        yield from iter_concedido_csv(path, chunk)


# --------------------------------------------------------------------------
# aggregation
# --------------------------------------------------------------------------
GRAIN = [
    "id_municipio",
    "ano",
    "mes",
    "especie_beneficio",
    "categoria_beneficio",
    "clientela",
    "sexo",
    "faixa_etaria",
]


class UnmappedLabelError(ValueError):
    """Raised when an espécie label has no code, rather than nulling it."""


def aggregate_concedido(
    path: Path, competencia_hint: int | None = None
) -> tuple[pd.DataFrame, dict]:
    """Aggregate one concedido file to the published grain.

    Returns the aggregated frame and a diagnostics dict recording how many
    rows lost their municipality or their age, which feeds the per-year
    coverage report.
    """
    esp_idx = especie_label_index()
    cat_idx = categoria_index()
    mun_idx = municipio_index()
    cells: dict[tuple, list[float]] = {}
    diag = {
        "linhas": 0,
        "sem_municipio": 0,
        "municipio_nao_encontrado": 0,
        "sem_idade": 0,
        "sem_competencia": 0,
        "nomes_nao_encontrados": {},
    }
    unmapped: set[str] = set()

    for batch in iter_concedido(path):
        for rec in batch:
            diag["linhas"] += 1
            comp = parse_competencia(rec["competencia"]) or competencia_hint
            if comp is None:
                diag["sem_competencia"] += 1
                continue
            code = rec["especie_codigo"]
            if code is None or (isinstance(code, str) and not code.strip()):
                lab = rec["especie_label"]
                code = esp_idx.get(expand_abbrev(lab)) or esp_idx.get(
                    norm_token(lab)
                )
                if code is None:
                    unmapped.add(str(lab))
                    continue
            code = int(code)

            uf, nome = parse_mun_resid(rec["mun_resid"])
            if uf is None:
                diag["sem_municipio"] += 1
                id_mun = None
            else:
                id_mun = lookup_municipio(mun_idx, uf, nome)
                if id_mun is None:
                    diag["municipio_nao_encontrado"] += 1
                    diag["nomes_nao_encontrados"][f"{uf}|{nome}"] = (
                        diag["nomes_nao_encontrados"].get(f"{uf}|{nome}", 0)
                        + 1
                    )

            idade = idade_em(parse_data(rec["dt_nascimento"]), comp)
            if idade is None:
                diag["sem_idade"] += 1
            key = (
                id_mun,
                comp // 100,
                comp % 100,
                code,
                cat_idx.get(code),
                clean_clientela(rec["clientela"]),
                clean_sexo(rec["sexo"]),
                faixa_etaria(idade),
            )
            rmi = parse_decimal(rec["qt_sm_rmi"])
            acc = cells.setdefault(key, [0, 0.0])
            acc[0] += 1
            if rmi is not None:
                acc[1] += rmi

    if unmapped:
        raise UnmappedLabelError(
            f"{path.name}: espécie labels with no code: {sorted(unmapped)}. "
            "Add them to constants.ESPECIE_LABEL_ALIAS."
        )

    df = pd.DataFrame(
        [(*k, v[0], v[1]) for k, v in cells.items()],
        columns=[*GRAIN, "quantidade", "valor_total_salarios_minimos"],
    )
    if not df.empty:
        sm = df["ano"] * 100 + df["mes"]
        df["valor_total"] = (
            df["valor_total_salarios_minimos"] * sm.map(salario_minimo_for)
        ).round(2)
        df["valor_total_salarios_minimos"] = df[
            "valor_total_salarios_minimos"
        ].round(3)
        df = df.sort_values(GRAIN, na_position="last").reset_index(drop=True)
    return df, diag


# --------------------------------------------------------------------------
# dicionário de espécies
# --------------------------------------------------------------------------
def build_dicionario_especie() -> pd.DataFrame:
    """The espécie code table, as published, plus the reform-stable grouping."""
    validate_reference_tables()
    cat = categoria_index()
    nat = {
        c: n for n, codes in constants.NATUREZA.value.items() for c in codes
    }
    dupes = [c for n, codes in constants.NATUREZA.value.items() for c in codes]
    if len(dupes) != len(set(dupes)):
        raise ValueError("NATUREZA assigns a code to more than one nature")
    unknown = set(dupes) - set(constants.ESPECIE.value)
    if unknown:
        raise ValueError(
            f"NATUREZA references unknown espécies: {sorted(unknown)}"
        )
    rows = [
        {
            "especie_beneficio": str(code),
            "nome_especie": label,
            "nome_especie_anterior": constants.NOME_ANTERIOR.value.get(code),
            "categoria_beneficio": cat[code],
            "natureza_beneficio": nat.get(code, "previdenciaria"),
            "observacao": constants.OBSERVACAO.value.get(code),
        }
        for code, label in constants.ESPECIE.value.items()
    ]
    return (
        pd.DataFrame(rows)
        .sort_values("especie_beneficio", key=lambda s: s.astype(int))
        .reset_index(drop=True)
    )


# --------------------------------------------------------------------------
# GEX lookup — the key that makes benefícios mantidos geocodable
# --------------------------------------------------------------------------
def truncated_mun_key(gex: str, uf: str, nome: str) -> str:
    """The 20-character field that benefícios mantidos actually stores."""
    return f"{gex}-{uf}-{nome}"[:20]


def build_gex_lookup(
    observed: dict[tuple[str, str, str], str],
) -> pd.DataFrame:
    """Turn (GEX, UF, full name) -> id_municipio into a truncated-key lookup.

    Benefícios mantidos truncates ``Município`` to 20 characters, leaving only
    11 of the name — ambiguous for 6.9% of municipalities on the name alone.
    Including the GEX prefix makes the truncated key unique, so the lookup is
    built from benefícios concedidos, which publishes the GEX, the full name
    and enough rows to cover every municipality.
    """
    keys: dict[str, set[str]] = {}
    meta: dict[str, tuple[str, str, str]] = {}
    for (gex, uf, nome), id_mun in observed.items():
        k = truncated_mun_key(gex, uf, nome)
        keys.setdefault(k, set()).add(id_mun)
        meta[k] = (gex, uf, nome)
    rows = [
        {
            "chave_truncada": k,
            "id_municipio": next(iter(v)),
            "gex": meta[k][0],
            "sigla_uf": meta[k][1],
            "nome": meta[k][2],
        }
        for k, v in sorted(keys.items())
        if len(v) == 1
    ]
    ambiguous = {k: sorted(v) for k, v in keys.items() if len(v) > 1}
    if ambiguous:
        print(
            f"WARNING: {len(ambiguous)} ambiguous truncated keys dropped: "
            f"{list(ambiguous.items())[:5]}"
        )
    return pd.DataFrame(rows)


# --------------------------------------------------------------------------
# parquet output
# --------------------------------------------------------------------------
def write_partitioned(
    df: pd.DataFrame,
    outdir: Path,
    table: str,
    partition_cols: list[str] | None = None,
) -> None:
    """Write hive-partitioned, all-STRING Snappy parquet.

    Every column is cast to string on purpose: staging is all-STRING by house
    convention and the dbt model safe_casts each column back, so the parquet
    schema carries column order, not types. The cast goes through arrow rather
    than ``astype(str)`` because the latter renders NULL as the literal "nan",
    which safe_cast will not turn back into NULL, and it runs after the real
    dtypes are set so that an integer year serialises as "2017", not "2017.0".
    """
    partition_cols = partition_cols or ["ano"]
    if df.empty:
        return
    arch = pd.read_csv(constants.ARCHITECTURE_DIR.value / f"{table}.csv")
    order = [c for c in arch["name"].tolist()]
    missing = set(order) - set(df.columns)
    if missing:
        raise ValueError(
            f"{table}: frame is missing architecture columns {sorted(missing)}"
        )
    df = df[order]

    types = dict(zip(arch["name"], arch["bigquery_type"], strict=True))
    out = pd.DataFrame(index=df.index)
    for c in order:
        s = df[c]
        bq = types[c]
        if bq == "INT64":
            s = pd.to_numeric(s, errors="coerce").astype("Int64")
        elif bq == "FLOAT64":
            s = pd.to_numeric(s, errors="coerce").astype("Float64")
        out[c] = s
    tbl = pa.Table.from_pandas(out, preserve_index=False)
    tbl = tbl.cast(pa.schema([pa.field(c, pa.string()) for c in order]))
    pq.write_to_dataset(
        tbl,
        root_path=str(outdir / table),
        partition_cols=partition_cols,
        compression="snappy",
        existing_data_behavior="overwrite_or_ignore",
    )
