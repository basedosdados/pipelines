"""Generic reader for ABS time-series workbooks.

Every release in the ABS "Price indexes and inflation" topic ships the same
workbook layout: an ``Index`` sheet naming the catalogue and the table, plus
one or more ``Data*`` sheets carrying a metadata block (Unit, Frequency,
Series Start/End, Series ID) above dated observation rows. One parser reads
all of them; the per-release differences live in ``releases.py``.

No Prefect imports here. The one-shot onboarding bootstrap and the recurring
pipeline both import these functions.
"""

from __future__ import annotations

import datetime as dt
import os
import re
from pathlib import Path

# pyrefly: ignore [untyped-import]
import openpyxl

from pipelines.datasets.au_abs_prices_inflation.constants import constants

RELEASES = constants.RELEASES.value
HEADERS = constants.HEADERS.value

_META_LABELS = {
    "Unit",
    "Series Type",
    "Data Type",
    "Frequency",
    "Collection Month",
    "Series Start",
    "Series End",
    "No. Obs",
    "Series ID",
}

_CATALOGUE_RE = re.compile(r"^(\d{4}\.\d)")


# --------------------------------------------------------------------------- #
# Download
# --------------------------------------------------------------------------- #
def _get(url: str, session=None):
    import requests

    getter = session or requests
    resp = getter.get(url, headers=HEADERS, timeout=180)
    resp.raise_for_status()
    return resp


def landing_html(release: str, session=None) -> str:
    """Fetch the release's ``latest-release`` landing page."""
    return _get(
        constants.RELEASE_LANDING_URL.value.format(**RELEASES[release]),
        session,
    ).text


def resolve_release_slug(release: str, html: str) -> str:
    """Read the current release slug (e.g. ``jun-2026``) off the landing page.

    The xlsx download URLs are dated by reference period and each release runs
    on its own cadence and its own slug shape -- Total Value of Dwellings uses
    ``jun-quarter-2026`` where the others use ``jun-2026`` -- so the slug is
    read from the page rather than reconstructed from the calendar.
    """
    m = re.search(RELEASES[release]["slug_re"], html)
    if not m:
        raise RuntimeError(
            f"could not resolve the ABS release slug for {release} "
            f"from {constants.RELEASE_LANDING_URL.value.format(**RELEASES[release])}"
        )
    return m.group(1)


def resolve_release_files(release: str, html: str) -> list[str]:
    """List this release's time-series workbook stems, from the landing page.

    Reading the links rather than hardcoding stems means a table ABS adds is
    picked up on its own. The per-release ``file_re`` keeps only the
    time-series workbooks: it drops the data cubes, which are a different
    layout this parser cannot read, and (for the Wage Price Index) the
    ``63450Table2ato9a`` consolidations, whose series all appear in the
    individual table workbooks as well -- verified to lose no series.
    """
    pattern = re.compile(RELEASES[release]["file_re"])
    stems = []
    for href in re.findall(r'href="([^"]+\.xlsx)"', html):
        stem = os.path.basename(href)[: -len(".xlsx")]
        if pattern.match(stem) and stem not in stems:
            stems.append(stem)
    if not stems:
        raise RuntimeError(
            f"no time-series workbooks found on the {release} landing page"
        )
    return sorted(stems)


def download_release(
    release: str, out_dir: str, session=None
) -> tuple[str, list[str]]:
    """Download every time-series workbook for the current release.

    Returns ``(slug, paths)``. The slug is recorded as source metadata and is
    what tells a scheduled run that ABS has published a new period.
    """
    html = landing_html(release, session)
    slug = resolve_release_slug(release, html)
    dest = Path(out_dir) / release
    dest.mkdir(parents=True, exist_ok=True)

    paths = []
    for stem in resolve_release_files(release, html):
        path = dest / f"{stem}.xlsx"
        if not path.exists():
            url = constants.RELEASE_FILE_URL.value.format(
                path=RELEASES[release]["path"], slug=slug, file=stem
            )
            path.write_bytes(_get(url, session).content)
        paths.append(str(path))
    return slug, paths


# --------------------------------------------------------------------------- #
# Parse
# --------------------------------------------------------------------------- #
def split_description(desc: str) -> list[str]:
    """Split an ABS Data Item Description into its parts.

    Split on ``" ; "`` -- spaces on both sides -- never on a bare semicolon.
    Total Value of Dwellings has measures that contain one
    (``Value of dwelling stock; Owned by All Sectors``), so a bare-semicolon
    split silently turns a two-part description into three and mislabels every
    dwelling-stock series. ABS also leaves a trailing ``" ;"`` artifact on the
    last part, which is stripped here.
    """
    # pyrefly: ignore [unnecessary-type-conversion]
    parts = (p.strip().rstrip(";").strip() for p in str(desc).split(" ; "))
    return [p for p in parts if p]


def parse_ts_workbook(path: str) -> tuple[list[dict], list[dict]]:
    """Parse one ABS time-series workbook into series and observation rows.

    Series metadata comes from the per-sheet block above the data rather than
    from the ``Index`` sheet, because that block is the one place carrying the
    unit and frequency alongside the Series ID for every column -- and units
    vary *within* a single ABS table, so they cannot be read per file.
    """
    base = os.path.basename(path).replace(".xlsx", "")
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)

    catalogue, table = None, None
    if "Index" in wb.sheetnames:
        for row in wb["Index"].iter_rows(values_only=True):
            value = row[1] if len(row) > 1 else None
            if not isinstance(value, str):
                continue
            match = _CATALOGUE_RE.match(value)
            if catalogue is None and match:
                catalogue = match.group(1)
            if table is None and value.upper().startswith("TABLE"):
                table = value.strip()

    series_rows: list[dict] = []
    obs_rows: list[dict] = []
    for sheet in (s for s in wb.sheetnames if s.lower().startswith("data")):
        rows = list(wb[sheet].iter_rows(values_only=True))
        if not rows:
            continue
        label_at = {}
        for i, row in enumerate(rows[:15]):
            head = row[0]
            if (
                isinstance(head, str)
                and head.strip().rstrip(".") in _META_LABELS
            ):
                label_at[head.strip().rstrip(".")] = i
        if "Series ID" not in label_at:
            continue

        desc_row = rows[0]
        sid_row = rows[label_at["Series ID"]]
        data_start = max(label_at.values()) + 1

        def cell(label, j, _rows=rows, _at=label_at):
            i = _at.get(label)
            if i is None:
                return None
            value = _rows[i][j]
            if isinstance(value, dt.datetime):
                return value.date()
            return str(value).strip() if value is not None else None

        cols = [
            j for j in range(1, len(sid_row)) if sid_row[j] not in (None, "")
        ]
        for j in cols:
            series_rows.append(
                {
                    "series_id": str(sid_row[j]).strip(),
                    "description": str(desc_row[j]),
                    "unit": cell("Unit", j),
                    "frequency": cell("Frequency", j),
                    "series_type": cell("Series Type", j),
                    "source_catalogue": catalogue,
                    "source_table": table,
                    "source_file": base,
                }
            )
        for row in rows[data_start:]:
            date = row[0]
            if not isinstance(date, dt.datetime):
                continue
            for j in cols:
                value = row[j]
                if not isinstance(value, (int, float)) or isinstance(
                    value, bool
                ):
                    continue
                obs_rows.append(
                    {
                        "series_id": str(sid_row[j]).strip(),
                        "date": date.date(),
                        "value": float(value),
                    }
                )
    wb.close()
    return series_rows, obs_rows


# --------------------------------------------------------------------------- #
# Period helpers
# --------------------------------------------------------------------------- #
def quarter_of(date: dt.date) -> int:
    """Quarter number for an ABS quarterly observation.

    ABS labels a quarter by its final month, so the June quarter is month 6.
    """
    return date.month // 3


def financial_year_of(date: dt.date) -> str:
    """Australian financial year for an ABS annual observation, e.g. ``2025-26``.

    ABS dates a financial-year series at the June ending the year.
    """
    return f"{date.year - 1}-{date.year % 100:02d}"
