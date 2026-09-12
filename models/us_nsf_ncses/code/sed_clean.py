"""Clean the NCSES SED published data tables into Data Basis tables.

Source
------
The Survey of Earned Doctorates releases one publication per survey cycle
("Doctorate Recipients from U.S. Universities"; NSF 25-349 is the 2024 cycle).
Each publication ships ~96 numbered data tables as Excel workbooks, bundled in
one ZIP:

    https://ncses.nsf.gov/pubs/<pub_id>/assets/data-tables/<pub_id>-data-tables-tables-excels.zip

SED *microdata* is restricted-use and is deliberately not touched here: only the
published aggregate tables are in scope.

Layout
------
Every workbook follows the same NCSES house layout, which is what makes one
parser enough for all of them:

    row 1   Table 1-5
    row 2   <title>
    row 3   (Number)                 <- unit statement
    row 4   <header, row 1>          <- merged cells span column groups
    row 5   <header, row 2>          <- only when row 4 has merged spans
    ...     <stub> <values...>       <- stub indentation carries the hierarchy
    ...     a <footnote> / Note(s): / Source(s):

Two features of the workbooks are load-bearing and are read from the cell
styles rather than the values: the stub's **indent level**, which reconstructs
the field/characteristic hierarchy, and the header's **merged ranges**, which
reconstruct the column groups.

Output
------
Two tables, hive-partitioned by ``reference_year`` (the survey cycle):

* ``sed_data_table``  one row per published table: id, title, unit statement
* ``sed_estimate``    one row per published cell, fully long

Each cycle restates its own history, so a cycle is a self-contained vintage:
filter to the most recent ``reference_year`` for the current numbers rather than
summing across cycles.

Run
---
    python models/us_nsf_ncses/code/sed_clean.py
"""

from __future__ import annotations

import os
import re
import sys
import zipfile
from pathlib import Path

import openpyxl
import pyarrow as pa
import pyarrow.parquet as pq
from openpyxl.cell.rich_text import CellRichText

DATA_DIR = Path(
    os.environ.get(
        "NCSES_DATA_DIR", os.path.expanduser("~/Downloads/us_nsf_ncses_data")
    )
)
INPUT_DIR = DATA_DIR / "input"
OUTPUT_DIR = DATA_DIR / "output"
WORK_DIR = DATA_DIR / "work" / "sedx"

# Survey cycle -> NCSES publication id. One entry per onboarded cycle.
CYCLES = {2024: "nsf25349"}

TABLE_ID_RE = re.compile(r"^Table\s+(\d+)[\u2013-](\d+)")
FOOTNOTE_START = re.compile(
    r"^(note\(s\)|source\(s\)|notes:|source:|n/a\b)", re.I
)

UNIT_PATTERNS = [
    (re.compile(r"\(number\)|\bnumber\b", re.I), "number"),
    (re.compile(r"%|\(percent\)|\bpercent\b", re.I), "percent"),
    (re.compile(r"\bdollars\b|\bsalary\b|\bdebt\b", re.I), "dollars"),
    (re.compile(r"median years|years to degree", re.I), "median_years"),
    (re.compile(r"\bmean\b", re.I), "mean"),
    (re.compile(r"\bmedian\b", re.I), "median"),
]

TABLE_COLUMNS = [
    "reference_year",
    "table_id",
    "table_group",
    "table_title",
    "unit_statement",
    "publication_id",
    "source_file",
    "estimate_count",
]

ESTIMATE_COLUMNS = [
    "reference_year",
    "table_id",
    "year",
    "row_number",
    "column_number",
    "row_label",
    "row_path",
    "row_level",
    "column_label",
    "column_path",
    "unit",
    "value",
]


def cell_text(value) -> str:
    """Render a header or stub cell as text, dropping NCSES footnote markers.

    The marker on ``"All doctorate recipients<sup>a</sup>"`` is a superscript
    run inside the cell's rich text, so the workbook is read with
    ``rich_text=True`` and superscript runs are discarded. Guessing from the
    characters instead would truncate real labels — ``"Male"`` would lose its
    ``e`` and ``"Science and engineering"`` its ``g``.
    """
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, CellRichText):
        parts = []
        for block in value:
            if isinstance(block, str):
                parts.append(block)
            elif getattr(block.font, "vertAlign", None) != "superscript":
                parts.append(block.text)
        return "".join(parts).strip()
    return str(value).strip()


def parse_number(value):
    """Return a float for a data cell, or None for a suppression marker."""
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace(",", "").replace("$", "")
    if text in {
        "",
        "-",
        "\u2013",
        "\u2014",
        "na",
        "NA",
        "n/a",
        "D",
        "S",
        "(X)",
        "(S)",
        "(D)",
    }:
        return None
    try:
        return float(text)
    except ValueError:
        return None


def infer_unit(*candidates: str) -> str:
    """Pick a unit from the most specific label available.

    Candidates are tried in order, so a column header that says ``"(%)"`` beats
    the table's ``"(Number and percent)"`` statement.
    """
    for text in candidates:
        if not text:
            continue
        for pattern, unit in UNIT_PATTERNS:
            if pattern.search(text):
                return unit
    return ""


def year_or_none(text: str):
    """Return a four-digit year from a header or stub label, else None."""
    text = text.strip()
    if re.fullmatch(r"(19|20)\d{2}", text):
        return int(text)
    return None


MAX_HEADER_ROW = 8


def header_chains(ws, merged) -> tuple[list[list[str]], int]:
    """Return (one label chain per column, first data row index, 0-based).

    The header starts at row 4 and is as deep as its merged ranges reach: a
    table with a plain header has none and is one row deep, while tables such
    as 3-3 and 6-1 nest three rows (citizenship, then ethnicity, then race).
    Reading only two rows collapses every race column onto the same pair of
    labels, which is what made 1,804 key collisions in the 2024 cycle.

    Horizontal merges spread their label across the span. Vertical merges are
    left alone, so a column spanning the whole header contributes one label
    rather than repeating it at every level.
    """
    grid = [[c.value for c in row] for row in ws.iter_rows()]
    ncol = max(len(r) for r in grid[:MAX_HEADER_ROW]) if grid else 0

    # Grow the header block while a merged range that starts inside it reaches
    # further down; the title merge on row 2 is outside the block and ignored.
    last = 4
    while True:
        reach = max(
            (r.max_row for r in merged if 4 <= r.min_row <= last),
            default=last,
        )
        reach = min(reach, MAX_HEADER_ROW)
        if reach == last:
            break
        last = reach

    rows = []
    for source in grid[3:last]:
        labels = [cell_text(v) for v in source]
        rows.append(labels + [""] * (ncol - len(labels)))
    for r in merged:
        if not (4 <= r.min_row <= last and r.max_col > r.min_col):
            continue
        label = rows[r.min_row - 4][r.min_col - 1]
        for c in range(r.min_col, min(r.max_col, ncol)):
            rows[r.min_row - 4][c] = label

    chains = [[row[c] for row in rows if row[c]] for c in range(ncol)]
    return chains, last


def parse_workbook(path: Path, reference_year: int, publication_id: str):
    """Parse one NCSES data table workbook into table + estimate rows."""
    wb = openpyxl.load_workbook(path, data_only=True, rich_text=True)
    ws = wb[wb.sheetnames[0]]
    merged = list(ws.merged_cells.ranges)
    grid = list(ws.iter_rows())

    tid_match = TABLE_ID_RE.match(cell_text(grid[0][0].value))
    if not tid_match:
        wb.close()
        raise RuntimeError(f"{path.name}: no table id in cell A1")
    table_id = f"{tid_match.group(1)}-{tid_match.group(2)}"
    title = cell_text(grid[1][0].value)
    unit_statement = cell_text(grid[2][0].value).strip("()")

    chains, first_data = header_chains(ws, merged)

    estimates: list[dict] = []
    ancestors: dict[int, str] = {}
    section_label = ""
    # Some tables restate the column labels part way down the stub — table 4-4
    # switches from "Median debt (dollars)" to a Number/Percent pair under each
    # debt concept. Such a row has an empty stub and text in the value columns,
    # and it re-labels every column below it until the next one.
    overrides: dict[int, str] = {}

    for row_number, row in enumerate(grid[first_data:], start=first_data + 1):
        stub_cell = row[0]
        stub = cell_text(stub_cell.value)
        if FOOTNOTE_START.match(stub):
            break
        if not stub:
            restated = {
                i + 1: cell_text(c.value)
                for i, c in enumerate(row[1:])
                if parse_number(c.value) is None and cell_text(c.value)
            }
            if restated:
                overrides = restated
            continue
        level = int(stub_cell.alignment.indent or 0)
        label = stub
        ancestors = {k: v for k, v in ancestors.items() if k < level}
        ancestors[level] = label
        path_labels = [ancestors[k] for k in sorted(ancestors)]

        stub_year = year_or_none(label)
        values = [parse_number(c.value) for c in row[1:]]
        if all(v is None for v in values):
            # A stub row with no numbers is a section heading; it also carries
            # the unit for the rows beneath it (e.g. "All recipients (number)").
            section_label = label
            continue

        for offset, value in enumerate(values):
            if value is None:
                continue
            col = offset + 1
            chain = chains[col] if col < len(chains) else []
            # A header level that is just a year is the table's time axis, not
            # a column label: it becomes `year` and drops out of the chain.
            header_year = next(
                (year_or_none(part) for part in chain if year_or_none(part)),
                None,
            )
            labels = [part for part in chain if not year_or_none(part)]
            if overrides.get(col):
                labels = [*labels, overrides[col]]
            year = stub_year or header_year or reference_year
            column_label = labels[-1] if labels else ""
            estimates.append(
                {
                    "reference_year": str(reference_year),
                    "table_id": table_id,
                    "year": str(year),
                    "row_number": str(row_number),
                    "column_number": str(col + 1),
                    "row_label": label,
                    "row_path": " > ".join(path_labels),
                    "row_level": str(level),
                    "column_label": column_label,
                    "column_path": " > ".join(labels),
                    "unit": infer_unit(
                        *reversed(labels),
                        section_label,
                        unit_statement,
                    ),
                    "value": repr(value),
                }
            )
    wb.close()

    table_row = {
        "reference_year": str(reference_year),
        "table_id": table_id,
        "table_group": tid_match.group(1),
        "table_title": title,
        "unit_statement": unit_statement,
        "publication_id": publication_id,
        "source_file": path.name,
        "estimate_count": str(len(estimates)),
    }
    return table_row, estimates


def extract_cycle(reference_year: int, publication_id: str) -> Path:
    """Unpack a cycle's data-table ZIP into the working directory."""
    zip_path = INPUT_DIR / f"sed{reference_year}_xlsx.zip"
    out = WORK_DIR / publication_id
    out.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(zip_path) as zf:
        zf.extractall(out)
    return out


def write_partition(table: str, reference_year: int, columns, rows) -> int:
    """Write one cycle of one table as all-STRING snappy parquet."""
    if not rows:
        return 0
    schema = pa.schema([(c, pa.string()) for c in columns])
    arrays = [
        pa.array(
            [r.get(c) if r.get(c) != "" else None for r in rows],
            type=pa.string(),
        )
        for c in columns
    ]
    out = OUTPUT_DIR / table / f"reference_year={reference_year}"
    out.mkdir(parents=True, exist_ok=True)
    pq.write_table(
        pa.Table.from_arrays(arrays, schema=schema),
        out / "data.parquet",
        compression="snappy",
    )
    return len(rows)


def main() -> int:
    for reference_year, publication_id in sorted(CYCLES.items()):
        folder = extract_cycle(reference_year, publication_id)
        books = sorted(folder.glob(f"{publication_id}-tab*.xlsx"))
        if not books:
            raise SystemExit(f"no workbooks under {folder}")
        tables, estimates = [], []
        for book in books:
            table_row, rows = parse_workbook(
                book, reference_year, publication_id
            )
            tables.append(table_row)
            estimates.extend(rows)
        tables.sort(
            key=lambda r: (
                int(r["table_group"]),
                int(r["table_id"].split("-")[1]),
            )
        )
        n_tables = write_partition(
            "sed_data_table", reference_year, TABLE_COLUMNS, tables
        )
        n_est = write_partition(
            "sed_estimate", reference_year, ESTIMATE_COLUMNS, estimates
        )
        unresolved = sum(1 for e in estimates if not e["unit"])
        with_year = sum(
            1 for e in estimates if int(e["year"]) != reference_year
        )
        print(
            f"SED {reference_year} ({publication_id}): {n_tables} tables, "
            f"{n_est:,} estimates; {with_year:,} carry a data year other than "
            f"the cycle year; unit unresolved on {unresolved:,} "
            f"({unresolved / max(n_est, 1):.1%})"
        )
        empty = [t["table_id"] for t in tables if t["estimate_count"] == "0"]
        if empty:
            raise RuntimeError(f"tables parsed to zero rows: {empty}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
